"""Preenche o modelo .docx com os dados do formulário e insere as imagens."""
import io
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph as DocxParagraph
from PIL import Image, ImageOps

Image.MAX_IMAGE_PIXELS = 60_000_000

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MODELO = os.path.join(ROOT, "assets", "Modelo-DOC-Solicitacao-Prefeitura.docx")

# Marcador do modelo  ->  chave dos dados
CAMPOS = {
    "NDaRegioDoConselhoPo": "regiao",
    "NomeCompleto": "nome",
    "NDoCPF": "cpf",
    "EnderecoCompleto": "endereco",
    "NDoCEP": "cep",
    "Bairro": "bairro",
    "TelefoneWhatsapp": "telefone",
    "DescrevaSuaDemanda": "demanda",
    "ProtocoloDoProcessoN": "protocolo",
}

# Área útil da página A4 com margens de 1" => 8,27" - 2" = 6,27"
LARGURA_UTIL_POL = 6.27


# --------------------------------------------------------------------------
# Imagens
# --------------------------------------------------------------------------
def preparar_imagem(blob: bytes, max_px: int, qualidade: int):
    """Corrige orientação, achata transparência, redimensiona e recomprime.

    Retorna (bytes_jpeg, largura_px, altura_px).
    """
    img = Image.open(io.BytesIO(blob))
    img = ImageOps.exif_transpose(img)
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        fundo = Image.new("RGB", img.size, (255, 255, 255))
        fundo.paste(img, mask=img.split()[-1])
        img = fundo
    else:
        img = img.convert("RGB")
    img.thumbnail((max_px, max_px), Image.LANCZOS)
    saida = io.BytesIO()
    img.save(saida, "JPEG", quality=qualidade, optimize=True)
    return saida.getvalue(), img.width, img.height


def _ajustar(w_px: int, h_px: int, max_w: float, max_h: float):
    """Escala (em polegadas) mantendo proporção dentro da caixa max_w x max_h."""
    esc = min(max_w / w_px, max_h / h_px)
    return Inches(w_px * esc), Inches(h_px * esc)


# --------------------------------------------------------------------------
# Texto
# --------------------------------------------------------------------------
def _substituir(par, chave_marcador: str, valor: str) -> bool:
    """Troca {{Marcador}} mesmo que o Word o tenha quebrado em vários runs.

    O valor herda a formatação do run que contém '{{'.
    """
    runs = list(par.runs)
    texto = "".join(r.text for r in runs)
    m = re.search(r"\{\{\s*" + re.escape(chave_marcador) + r"\s*\}\}", texto)
    if not m:
        return False
    ini, fim = m.span()
    pos = 0
    primeiro = True
    for r in runs:
        t = r.text
        r_ini, r_fim = pos, pos + len(t)
        pos = r_fim
        if r_fim <= ini or r_ini >= fim:
            continue
        a = max(ini, r_ini) - r_ini
        b = min(fim, r_fim) - r_ini
        if primeiro:
            r.text = t[:a] + valor + t[b:]
            primeiro = False
        else:
            r.text = t[b:]
    return True


def _remover(par):
    par._p.getparent().remove(par._p)


def _paragrafo_com_marcador_imagem(doc, nome: str):
    for p in doc.paragraphs:
        for descr in p._p.xpath(".//wp:docPr/@descr"):
            if nome in descr:
                return p
    return None


def _limpar_imagem_modelo(doc, par):
    """Remove a imagem de exemplo do parágrafo e descarta o arquivo do pacote."""
    rids = par._p.xpath(".//a:blip/@r:embed")
    for r in list(par.runs):
        if r._r.xpath(".//w:drawing"):
            r._r.getparent().remove(r._r)
    for rid in rids:
        try:
            doc.part.drop_rel(rid)
        except KeyError:
            pass
    for r in list(par.runs):  # tira espaços soltos do modelo
        if not r.text.strip():
            r._r.getparent().remove(r._r)


# --------------------------------------------------------------------------
# Fotos da demanda (grade sem bordas)
# --------------------------------------------------------------------------
def _sem_quebra_de_linha(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _tabela_fotos(doc, ancora, fotos):
    n = len(fotos)
    colunas = 1 if n == 1 else 2
    linhas = (n + colunas - 1) // colunas
    larg_cel = LARGURA_UTIL_POL / colunas
    caixa = (5.8, 4.4) if n == 1 else (larg_cel - 0.15, 3.4)

    tbl = doc.add_table(rows=linhas, cols=colunas)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for col in tbl.columns:
        col.width = Inches(larg_cel)

    for i, (dados, w, h) in enumerate(fotos):
        cel = tbl.rows[i // colunas].cells[i % colunas]
        cel.width = Inches(larg_cel)
        par = cel.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(6)
        lw, lh = _ajustar(w, h, *caixa)
        par.add_run().add_picture(io.BytesIO(dados), width=lw, height=lh)

    for row in tbl.rows:
        _sem_quebra_de_linha(row)
        for cel in row.cells:  # células vazias da última linha
            cel.width = Inches(larg_cel)
            if not cel.paragraphs[0].runs:
                cel.paragraphs[0].paragraph_format.space_after = Pt(0)

    ancora._p.addnext(tbl._tbl)


# --------------------------------------------------------------------------
# Principal
# --------------------------------------------------------------------------
def gerar_docx(dados: dict, fotos: list, print_app=None) -> bytes:
    """
    dados     : dict já validado (regiao, nome, cpf, ...)
    fotos     : lista de (bytes_jpeg, w_px, h_px)  - obrigatório (>=1)
    print_app : (bytes_jpeg, w_px, h_px) ou None
    """
    doc = Document(MODELO)

    # 1) Textos -------------------------------------------------------------
    for par in list(doc.paragraphs):
        for marcador, chave in CAMPOS.items():
            if f"{marcador}" in "".join(r.text for r in par.runs):
                _substituir(par, marcador, dados.get(chave, ""))

    # 2) Fotos da demanda ---------------------------------------------------
    par_fotos = _paragrafo_com_marcador_imagem(doc, "demandaImagens")
    if par_fotos is not None:
        _limpar_imagem_modelo(doc, par_fotos)
        _tabela_fotos(doc, par_fotos, fotos)
        _remover(par_fotos)

    # 3) Print do aplicativo (opcional) --------------------------------------
    par_print = _paragrafo_com_marcador_imagem(doc, "solicitacaoPrint")
    if par_print is not None:
        rotulo = par_print._p.getprevious()
        espaco = par_print._p.getnext()  # parágrafo vazio entre print e protocolo
        _limpar_imagem_modelo(doc, par_print)
        if print_app:
            if rotulo is not None:  # mantém "Solicitação:" junto da imagem
                DocxParagraph(rotulo, par_print._parent).paragraph_format.keep_with_next = True
            dados_img, w, h = print_app
            lw, lh = _ajustar(w, h, 3.2, 4.6)
            par_print.add_run().add_picture(io.BytesIO(dados_img), width=lw, height=lh)
        else:
            for el in (rotulo, par_print._p, espaco):
                if el is not None and el.getparent() is not None:
                    if el is espaco and "".join(el.itertext()).strip():
                        continue  # não apaga se tiver texto (ex.: Protocolo)
                    el.getparent().remove(el)

    # 4) Protocolo (opcional) ------------------------------------------------
    if not dados.get("protocolo"):
        for par in list(doc.paragraphs):
            if par.text.strip().startswith("Protocolo"):
                _remover(par)

    doc.core_properties.title = f"Solicitação - {dados.get('nome', '')}"
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()
